# -*- coding: utf-8 -*-
"""针对 9 道问题题目的定向测试脚本。

测试修复后的效果，输出格式与 result_3.xlsx 一致。
输出: result/test_problem.xlsx
"""
import json
import re
import time
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.graph import build_graph
from tool.visualizer import set_chart_prefix

XLSX_PATH = Path("D:/td26/B题测试数据/测试数据/附件6：问题汇总.xlsx")
OUT_PATH = Path("result/test_problem.xlsx")

PROBLEM_IDS = {"B4007", "B4008", "B4013", "B4015", "B4037", "B4041", "B4045", "B4046", "B4055"}

_SQL_RE = re.compile(r"SQL:\s*(.+?)(?:\n\n结果:|\Z)", re.DOTALL)


def _extract_sql(sql_result: str) -> str:
    sqls = _SQL_RE.findall(sql_result)
    return "\n".join(s.strip() for s in sqls if s.strip()) if sqls else "无"


def load_questions() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    questions = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        eid, qtype, raw = row
        if not raw or eid not in PROBLEM_IDS:
            continue
        try:
            turns = json.loads(raw)
        except json.JSONDecodeError:
            continue
        questions.append({"id": eid, "type": qtype, "turns": turns})
    return questions


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    questions = load_questions()
    print(f"待测试 {len(questions)} 道问题题目: {[q['id'] for q in questions]}")

    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "问题题目测试"
    ws.append(["编号", "问题", "SQL查询语法", "回答"])

    graph = build_graph()
    wrap_align = Alignment(wrap_text=True, vertical="top")

    for i, q in enumerate(questions, 1):
        thread_id = i + 1000
        config = {"configurable": {"thread_id": str(thread_id)}}
        set_chart_prefix(q["id"])

        all_answers = []
        all_sql = []
        t0 = time.time()

        for turn in q["turns"]:
            query = turn["Q"]
            try:
                state = graph.invoke(
                    {"messages": [{"role": "user", "content": query}]},
                    config=config,
                )
                formatted = state.get("formatted_answer", "")
                if formatted:
                    all_answers.append(json.loads(formatted))
                sql_result = state.get("sql_result", "")
                if sql_result:
                    all_sql.append(_extract_sql(sql_result))

                # 打印诊断信息
                rag_result = state.get("rag_result", "")
                sub_tasks = state.get("sub_tasks", [])
                tools = [t.get("tool", "") for t in sub_tasks]
                print(f"  轮次: {query[:50]}...")
                print(f"  工具: {tools}")
                if rag_result:
                    try:
                        robj = json.loads(rag_result)
                        refs = robj.get("references", [])
                        content = robj.get("content", "")[:100]
                        print(f"  RAG引用数: {len(refs)}, 内容: {content}...")
                    except:
                        pass
                if sql_result:
                    print(f"  SQL结果: {sql_result[:100]}...")
            except Exception as e:
                all_answers.append({"Q": query, "A": {"content": f"执行失败: {e}"}})
                print(f"  错误: {e}")

        elapsed = time.time() - t0
        sql_str = "\n".join(all_sql) if all_sql else "无"
        answer_str = json.dumps(all_answers, ensure_ascii=False, indent=2)

        row = [q["id"], json.dumps(q["turns"], ensure_ascii=False), sql_str, answer_str]
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = wrap_align

        wb_out.save(OUT_PATH)
        print(f"[{i}/{len(questions)}] {q['id']} ({q['type']}) {elapsed:.1f}s OK\n")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 80
    wb_out.save(OUT_PATH)
    print(f"\n完成，结果写入 {OUT_PATH}")


if __name__ == "__main__":
    main()
