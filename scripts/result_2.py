# -*- coding: utf-8 -*-
"""批量测试脚本：读取附件4问题集，逐题运行 Agent，输出结果到 Excel。

用法:
    cd FS
    uv run python scripts/result_2.py --input ../正式数据/附件4：问题汇总.xlsx --output result/result_2.xlsx
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import sys
import time
import uuid
import warnings
from pathlib import Path

import pandas as pd
from langchain_core.messages import HumanMessage
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Fix Windows console encoding
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

# ── 静默噪音 ──
warnings.filterwarnings("ignore")
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")
os.environ.setdefault("MODELSCOPE_LOG_LEVEL", str(logging.ERROR))
os.environ.setdefault("TQDM_DISABLE", "1")
for _name in ("modelscope", "sentence_transformers", "transformers",
              "torch", "jieba", "chromadb", "httpx", "urllib3"):
    logging.getLogger(_name).setLevel(logging.ERROR)
logging.basicConfig(level=logging.WARNING, force=True)

# 设置项目路径
_PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_PROJECT_ROOT))

import jieba
jieba.setLogLevel(logging.ERROR)

from agent.graph import build_graph, _execute_sql, _execute_rag, _execute_viz
from agent.planner import _parse_tasks, _get_planner_llm, PLANNER_SYSTEM_PROMPT
from tool.visualizer import set_chart_prefix
from langchain_core.messages import SystemMessage


def parse_questions(xlsx_path: str) -> list[dict]:
    """解析问题汇总 Excel，返回 [{编号, 问题类型, questions: [str, ...]}]"""
    df = pd.read_excel(xlsx_path)
    rows = []
    for _, row in df.iterrows():
        qid = str(row.iloc[0])
        qtype = str(row.iloc[1])
        raw = str(row.iloc[2])
        try:
            qs = json.loads(raw)
            questions = [q["Q"] for q in qs]
        except (json.JSONDecodeError, KeyError):
            questions = [raw]
        rows.append({"id": qid, "type": qtype, "questions": questions})
    return rows


def run_single_session(graph, questions: list[str], qid: str) -> dict:
    """运行一组多轮对话问题，返回详细结果。"""
    thread_id = f"test_{qid}_{uuid.uuid4().hex[:8]}"
    config = {"configurable": {"thread_id": thread_id}}

    all_turns = []
    all_logs = []

    for turn_idx, q in enumerate(questions):
        turn_start = time.time()
        log_lines = [f"  [Turn {turn_idx+1}] Q: {q}"]

        try:
            result = graph.invoke(
                {"messages": [HumanMessage(content=q)]},
                config=config,
            )

            # 提取各阶段信息
            sub_tasks = result.get("sub_tasks", [])
            sql_result = result.get("sql_result", "")
            rag_result = result.get("rag_result", "")
            plot_path = result.get("plot_path", "")
            final_answer = result.get("final_answer", "")
            formatted_answer = result.get("formatted_answer", "")

            # 打印 Planner 结果
            tools = [t.get("tool", "") for t in sub_tasks]
            log_lines.append(f"  [Planner] tasks={len(sub_tasks)}, tools={tools}")
            if sub_tasks:
                for t in sub_tasks:
                    log_lines.append(f"    SubTask: id={t.get('id')}, tool={t.get('tool')}, query={t.get('query')}, depends={t.get('depends_on')}")

            # 打印 SQL 过程
            if sql_result:
                sql_preview = sql_result[:500].replace('\n', '\n    ')
                log_lines.append(f"  [SQL Result]\n    {sql_preview}")

            # 打印 RAG 过程
            if rag_result:
                try:
                    rag_obj = json.loads(rag_result)
                    rag_content = rag_obj.get("content", "")[:300]
                    refs = rag_obj.get("references", [])
                    log_lines.append(f"  [RAG Content] {rag_content}...")
                    log_lines.append(f"  [RAG Refs] {len(refs)} references")
                    for ref in refs:
                        log_lines.append(f"    - {ref.get('paper_path', 'N/A')}: {ref.get('text', '')[:80]}...")
                except (json.JSONDecodeError, TypeError):
                    log_lines.append(f"  [RAG Raw] {rag_result[:200]}")

            # 打印可视化
            if plot_path:
                log_lines.append(f"  [Plot] {plot_path}")

            # 打印最终回答
            answer_preview = final_answer[:200] if final_answer else "(empty)"
            log_lines.append(f"  [Answer] {answer_preview}...")

            elapsed = time.time() - turn_start
            log_lines.append(f"  [Time] {elapsed:.1f}s")

            # 解析 formatted_answer
            turn_data = {}
            if formatted_answer:
                try:
                    turn_data = json.loads(formatted_answer)
                except (json.JSONDecodeError, TypeError):
                    turn_data = {"Q": q, "A": {"content": final_answer}}
            else:
                turn_data = {"Q": q, "A": {"content": final_answer}}

            import re as _re
            sql_stmts = []
            for t in sub_tasks:
                if t.get("tool") == "sql" and t.get("result"):
                    for m in _re.finditer(r"(?:SQL(?:\(分解执行\))?:\s*\n?)((?:SELECT|--).+?)(?:\n\n结果|\n\n---|\Z)", t["result"], _re.DOTALL):
                        sql_stmts.append(m.group(1).strip())
            turn_data["_sql"] = "\n".join(sql_stmts)

            chart_kw_map = {"折线": "line", "柱状": "bar", "饼": "pie", "雷达": "radar",
                            "直方": "histogram", "双条形": "double_bar", "水平": "double_bar",
                            "散点": "scatter", "表格": "table", "箱线": "boxplot"}
            detected_chart = ""
            for t in sub_tasks:
                if t.get("tool") == "visualize" and t.get("result", "").endswith(".png"):
                    vq = t.get("query", "")
                    for kw, ct in chart_kw_map.items():
                        if kw in vq:
                            detected_chart = ct
                            break
                    if not detected_chart:
                        detected_chart = "chart"
            turn_data["_chart_type"] = detected_chart

            all_turns.append(turn_data)

        except Exception as e:
            log_lines.append(f"  [ERROR] {e}")
            all_turns.append({"Q": q, "A": {"content": f"执行错误: {e}"}})
            elapsed = time.time() - turn_start
            log_lines.append(f"  [Time] {elapsed:.1f}s")

        log_text = "\n".join(log_lines)
        all_logs.append(log_text)
        try:
            print(log_text)
        except UnicodeEncodeError:
            print(log_text.encode('utf-8', errors='replace').decode('utf-8', errors='replace'))

    return {"turns": all_turns, "logs": all_logs}


def write_results_xlsx(results: list[dict], output_path: str):
    """将结果写入 Excel 文件。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "测试结果"

    headers = ["编号", "问题", "SQL 查询语句", "图形格式", "回答"]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row = 2
    for r in results:
        qid = r["id"]
        turns = r["result"]["turns"]

        q_list = [{"Q": t.get("Q", "")} for t in turns]
        q_text = json.dumps(q_list, ensure_ascii=False)

        sql_parts = []
        chart_types = set()
        for t in turns:
            sql = t.get("_sql", "")
            if sql:
                sql_parts.append(sql)
            ct = t.get("_chart_type", "")
            if ct:
                chart_types.add(ct)
        sql_text = "\n---\n".join(sql_parts) if sql_parts else ""

        chart_map = {
            "line": "折线图", "bar": "柱状图", "pie": "饼图",
            "radar": "雷达图", "histogram": "直方图", "double_bar": "双条形图",
            "scatter": "散点图", "table": "表格", "boxplot": "箱线图",
        }
        chart_text = "、".join(chart_map.get(c, c) for c in chart_types) if chart_types else "无"

        clean_turns = []
        for t in turns:
            ct = {k: v for k, v in t.items() if not k.startswith("_")}
            clean_turns.append(ct)
        answer_json = json.dumps(clean_turns, ensure_ascii=False)

        ws.cell(row=row, column=1, value=qid)
        ws.cell(row=row, column=2, value=q_text)
        cell_sql = ws.cell(row=row, column=3, value=sql_text)
        cell_sql.alignment = Alignment(wrap_text=True)
        ws.cell(row=row, column=4, value=chart_text)
        cell_a = ws.cell(row=row, column=5, value=answer_json)
        cell_a.alignment = Alignment(wrap_text=True)
        row += 1

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 80

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    try:
        wb.save(output_path)
        print(f"\n结果已保存: {output_path}")
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        fallback = f"{base}_{uuid.uuid4().hex[:4]}{ext}"
        wb.save(fallback)
        print(f"\n原路径被占用，结果已保存: {fallback}")


def main():
    parser = argparse.ArgumentParser(description="批量测试 FS Agent")
    parser.add_argument("--input", required=True, help="问题汇总 xlsx 路径")
    parser.add_argument("--output", required=True, help="输出结果 xlsx 路径")
    parser.add_argument("--start", type=int, default=0, help="从第N题开始（0-indexed）")
    parser.add_argument("--limit", type=int, default=0, help="最多测试N题（0=全部）")
    args = parser.parse_args()

    print(f"加载问题集: {args.input}")
    questions = parse_questions(args.input)
    print(f"共 {len(questions)} 题")

    if args.start > 0:
        questions = questions[args.start:]
        print(f"从第 {args.start} 题开始")
    if args.limit > 0:
        questions = questions[:args.limit]
        print(f"限制测试 {args.limit} 题")

    print("构建 Agent Graph...")
    graph = build_graph()
    print("Agent 就绪\n")

    results = []
    total = len(questions)
    for i, q in enumerate(questions):
        print(f"\n{'='*60}")
        print(f"[{i+1}/{total}] {q['id']} ({q['type']})")
        print(f"  Questions: {q['questions']}")
        print(f"{'='*60}")

        set_chart_prefix(q["id"])
        result = run_single_session(graph, q["questions"], q["id"])
        results.append({"id": q["id"], "type": q["type"], "result": result})

        # 增量保存，防止中途崩溃丢失结果
        if (i + 1) % 5 == 0 or i == total - 1:
            write_results_xlsx(results, args.output)

    print(f"\n全部完成！共测试 {total} 题。")


if __name__ == "__main__":
    main()
