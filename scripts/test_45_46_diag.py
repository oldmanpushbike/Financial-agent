# -*- coding: utf-8 -*-
"""单独测试 B4045 和 B4046 — 增加 planner 诊断。"""
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.graph import build_graph
from agent.planner import _get_planner_llm, PLANNER_SYSTEM_PROMPT
from langchain_core.messages import HumanMessage, SystemMessage
from tool.visualizer import set_chart_prefix

XLSX_PATH = Path("D:/td26/B题测试数据/测试数据/附件6：问题汇总.xlsx")

import openpyxl

def load_target_questions(ids):
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    questions = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        eid, qtype, raw = row
        if not raw or eid not in ids:
            continue
        try:
            turns = json.loads(raw)
        except json.JSONDecodeError:
            continue
        questions.append({"id": eid, "type": qtype, "turns": turns})
    return questions

def test_planner_only(query):
    llm = _get_planner_llm()
    resp = llm.invoke([SystemMessage(content=PLANNER_SYSTEM_PROMPT), HumanMessage(content=query)])
    print(f"  Planner raw: {resp.content[:500]}")
    return resp.content

def main():
    target_ids = {"B4045", "B4046"}
    questions = load_target_questions(target_ids)
    print(f"测试 {len(questions)} 道: {[q['id'] for q in questions]}")

    for q in questions:
        print(f"\n{'='*60}")
        print(f"[{q['id']}]")
        for turn in q["turns"]:
            query = turn["Q"]
            print(f"  Q: {query}")
            print(f"  --- Planner 单独测试 ---")
            test_planner_only(query)

if __name__ == "__main__":
    main()
