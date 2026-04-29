# -*- coding: utf-8 -*-
"""批量测试附件6的80道题，按提交格式输出结果。

输出: result/result_3.xlsx
图片: result/img/B2001_1.png, B2001_2.png, ...
"""
import json
import re
import time
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.graph import build_graph
from tool.visualizer import set_chart_prefix

XLSX_PATH = Path("D:/td26/B题测试数据/测试数据/附件6：问题汇总.xlsx")
OUT_PATH = Path("result/result_3.xlsx")

_SQL_RE = re.compile(r"SQL:\s*(.+?)(?:\n\n结果:|\Z)", re.DOTALL)


def _extract_sql(sql_result: str) -> str:
    """从 sql_result 中提取所有 SQL 语句。"""
    sqls = _SQL_RE.findall(sql_result)
    return "\n".join(s.strip() for s in sqls if s.strip()) if sqls else "无"


def load_questions() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    questions = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        eid, qtype, raw = row
        if not raw:
            continue
        try:
            turns = json.loads(raw)
        except json.JSONDecodeError:
            continue
        questions.append({"id": eid, "type": qtype, "turns": turns})
    return questions


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    questions = load_questions()
    print(f"共 {len(questions)} 道题")

    # 断点续跑：读取已有结果，跳过已完成的题目
    done_ids = set()
    if OUT_PATH.exists():
        wb_exist = openpyxl.load_workbook(OUT_PATH)
        ws_exist = wb_exist.active
        for row in ws_exist.iter_rows(min_row=2, values_only=True):
            if row[0]:
                done_ids.add(row[0])
        wb_exist.close()
        wb_out = openpyxl.load_workbook(OUT_PATH)
        ws = wb_out.active
        print(f"已完成 {len(done_ids)} 题，续跑剩余题目")
    else:
        wb_out = openpyxl.Workbook()
        ws = wb_out.active
        ws.title = "附件6结果"
        ws.append(["编号", "问题", "SQL查询语法", "回答"])

    graph = build_graph()
    thread_id = len(done_ids)
    wrap_align = Alignment(wrap_text=True, vertical="top")

    for i, q in enumerate(questions, 1):
        if q["id"] in done_ids:
            continue
        thread_id += 1
        config = {"configurable": {"thread_id": str(thread_id)}}
        set_chart_prefix(q["id"])

        all_answers = []
        all_sql = []
        t0 = time.time()

        for turn in q["turns"]:
            query = turn["Q"]
            try:
                state = graph.invoke(
                    {"messages": [{"role": "user", "content": query}]},
                    config=config,
                )
                formatted = state.get("formatted_answer", "")
                if formatted:
                    all_answers.append(json.loads(formatted))
                sql_result = state.get("sql_result", "")
                if sql_result:
                    all_sql.append(_extract_sql(sql_result))
            except Exception as e:
                all_answers.append({"Q": query, "A": {"content": f"执行失败: {e}"}})

        elapsed = time.time() - t0
        sql_str = "\n".join(all_sql) if all_sql else "无"
        answer_str = json.dumps(all_answers, ensure_ascii=False, indent=2)

        row = [q["id"], json.dumps(q["turns"], ensure_ascii=False), sql_str, answer_str]
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = wrap_align

        # 每题写一次，防止中途崩溃丢数据
        wb_out.save(OUT_PATH)
        print(f"[{i}/{len(questions)}] {q['id']} ({q['type']}) {len(q['turns'])}轮 {elapsed:.1f}s")

    # 调整列宽
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 80
    wb_out.save(OUT_PATH)
    print(f"\n完成，结果写入 {OUT_PATH}")


if __name__ == "__main__":
    main()
