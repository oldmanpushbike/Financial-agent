# -*- coding: utf-8 -*-
"""单独测试 B4045 和 B4046 的完整 pipeline。"""
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.graph import build_graph
from tool.visualizer import set_chart_prefix

XLSX_PATH = Path("D:/td26/B题测试数据/测试数据/附件6：问题汇总.xlsx")

import openpyxl

def load_target_questions(ids):
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    questions = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        eid, qtype, raw = row
        if not raw or eid not in ids:
            continue
        try:
            turns = json.loads(raw)
        except json.JSONDecodeError:
            continue
        questions.append({"id": eid, "type": qtype, "turns": turns})
    return questions

def main():
    target_ids = {"B4045", "B4046"}
    questions = load_target_questions(target_ids)
    print(f"测试 {len(questions)} 道: {[q['id'] for q in questions]}")

    graph = build_graph()

    for i, q in enumerate(questions, 1):
        thread_id = 9000 + i
        config = {"configurable": {"thread_id": str(thread_id)}}
        set_chart_prefix(q["id"])

        print(f"\n{'='*60}")
        print(f"[{q['id']}] 类型: {q['type']}")

        t0 = time.time()
        for turn in q["turns"]:
            query = turn["Q"]
            print(f"  Q: {query}")
            try:
                state = graph.invoke(
                    {"messages": [{"role": "user", "content": query}]},
                    config=config,
                )
                formatted = state.get("formatted_answer", "")
                sql_result = state.get("sql_result", "")
                rag_result = state.get("rag_result", "")
                sub_tasks = state.get("sub_tasks", [])
                tools = [t.get("tool", "") for t in sub_tasks] if sub_tasks else []

                print(f"  Tools: {tools}")
                if rag_result:
                    try:
                        robj = json.loads(rag_result)
                        refs = robj.get("references", [])
                        content = robj.get("content", "")[:200]
                        print(f"  RAG refs: {len(refs)}, content: {content}")
                    except:
                        print(f"  RAG raw: {rag_result[:200]}")
                if sql_result:
                    print(f"  SQL: {sql_result[:200]}")
                if formatted:
                    print(f"  Answer: {formatted[:300]}")
                else:
                    print(f"  Answer: (empty)")
                    final = state.get("final_answer", "")
                    if final:
                        print(f"  Final: {final[:300]}")
            except Exception as e:
                print(f"  ERROR: {e}")
                import traceback
                traceback.print_exc()

        elapsed = time.time() - t0
        print(f"  Time: {elapsed:.1f}s")

if __name__ == "__main__":
    main()
